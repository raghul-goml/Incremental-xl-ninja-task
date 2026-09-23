import sys
import asyncio
import io

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

from fastapi import (
    FastAPI,
    UploadFile,
    File,
    Depends,
    HTTPException
)

from sqlalchemy.orm import Session

from openpyxl import load_workbook

from .database import Base, engine, get_db
from .models import ExcelRow
from .crawlAI import crawl_url

from pydantic import BaseModel, Field

from typing import Optional


app = FastAPI(
    title="Incremental Excel Upload API"
)


Base.metadata.create_all(bind=engine)


@app.get("/")
def root():
    return {
        "message": "Incremental Excel API is running"
    }


@app.get("/users")
def get_user(name):
    return {
        "message": f"Hello, {name}!"
    }


@app.get("/crawl")
@app.post("/crawl")
async def run_crawl():
    result = await crawl_url()
    if not result["success"]:
        raise HTTPException(
            status_code=500,
            detail=f"Crawl failed: {result['error']}"
        )
    return result


@app.post("/upload")
async def upload_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only XLSX files are supported"
        )

    contents = await file.read()

    workbook = load_workbook(
        filename=io.BytesIO(contents),
        data_only=True
    )

    sheet = workbook.active

    headers = [
        cell.value
        for cell in sheet[1]
    ]

    processed_rows = []
    updated_rows = []
    skipped_rows = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        row_data = dict(zip(headers, row))

        if row_data.get("ID") is None:
            continue

        row_id = str(row_data["ID"])

        existing_row = (
            db.query(ExcelRow)
            .filter(
                ExcelRow.row_id == row_id
            )
            .first()
        )

        if not existing_row:

            new_row = ExcelRow(
                row_id=row_id,
                name=row_data.get("Name"),
                service=row_data.get("Service"),
                status=row_data.get("Status")
            )

            db.add(new_row)

            processed_rows.append(row_data)

        else:

            changes = {}

            if existing_row.name != row_data.get("Name"):
                changes["Name"] = {
                    "old": existing_row.name,
                    "new": row_data.get("Name")
                }
                existing_row.name = row_data.get("Name")

            if existing_row.service != row_data.get("Service"):
                changes["Service"] = {
                    "old": existing_row.service,
                    "new": row_data.get("Service")
                }
                existing_row.service = row_data.get("Service")

            if existing_row.status != row_data.get("Status"):
                changes["Status"] = {
                    "old": existing_row.status,
                    "new": row_data.get("Status")
                }
                existing_row.status = row_data.get("Status")

            if changes:

                updated_rows.append({
                    "row_id": row_id,
                    "changes": changes
                })

            else:

                skipped_rows.append(row_id)

    db.commit()

    return {
        "filename": file.filename,
        "processed_count": len(processed_rows),
        "updated_count": len(updated_rows),
        "skipped_count": len(skipped_rows),
        "processed_rows": processed_rows,
        "updated_rows": updated_rows,
        "skipped_rows": skipped_rows
    }